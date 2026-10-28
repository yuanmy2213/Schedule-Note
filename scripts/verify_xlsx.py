"""Independent consumer check for the workbook emitted by the C# tests."""
from pathlib import Path
import json
import openpyxl

path = Path('artifacts/verification.xlsx')
book = openpyxl.load_workbook(path, data_only=False)
assert book.sheetnames == ['事项明细', '每日汇总']
sheet = book['事项明细']
assert sheet['B2'].value == '独立读取校验 😀'
assert sheet['D2'].value is None
assert sheet['E2'].value.hour == 17
assert sheet['J2'].value == 42
assert sheet.freeze_panes == 'A2'
assert book['每日汇总']['F2'].value == .5
for row in sheet:
    for cell in row:
        assert cell.data_type != 'f', 'Untrusted user text became a formula'
Path('artifacts/test-results').mkdir(parents=True, exist_ok=True)
Path('artifacts/test-results/xlsx-independent.json').write_text(json.dumps({
    'passed': True, 'consumer': 'openpyxl', 'sheets': book.sheetnames,
    'checks': ['unicode', 'nullable-times', 'datetime-cells', 'duration', 'summary', 'freeze-row', 'no-formulas'],
}, ensure_ascii=False, indent=2), encoding='utf-8')
print('Independent workbook consumer validation passed.')
